import openpyxl
import os
import json
import datetime
import pandas as pd

file_path = os.path.join(os.path.dirname(__file__), "Controle de Estudos2.xlsm")
wb = openpyxl.load_workbook(file_path, data_only=True)

all_records = []

for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        continue
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    print(f"Sheet: {sheetname}, Headers: {headers}")
    
    for r_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        disc = row[0]
        dia = row[1]
        h_ini = row[2]
        h_fim = row[3]
        tempo = row[4]
        
        if disc is None and dia is None:
            continue
            
        all_records.append({
            "fase": sheetname,
            "disciplina": str(disc).strip() if disc else "Não especificado",
            "dia": str(dia) if dia else None,
            "hora_inicio": str(h_ini) if h_ini else None,
            "hora_fim": str(h_fim) if h_fim else None,
            "tempo_estudo": str(tempo) if tempo else None,
            "raw_tempo": repr(tempo)
        })

print(f"Total registros lidos: {len(all_records)}")
df = pd.DataFrame(all_records)
print(df.head(10))
print("\nDisciplinas únicas:")
print(df["disciplina"].value_counts())
