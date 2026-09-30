import os
import json
import datetime
import openpyxl
import math

file_path = os.path.join(os.path.dirname(__file__), "Controle de Estudos2.xlsm")
wb = openpyxl.load_workbook(file_path, data_only=True)

CATEGORIAS = {
    # SQL
    "SQL Para Análise de Dados e Data Science": "SQL & Dados",
    "Introdução ao SQL": "SQL & Dados",
    "Filtrando uma query e colocando condicionais": "SQL & Dados",
    "SQL a arte de agregar": "SQL & Dados",
    
    # R & Python
    "Manipulação de dados Tidyverse": "Programação (R / Python)",
    "Inicializando o R": "Programação (R / Python)",
    "Introdução ao R": "Programação (R / Python)",
    "Trabalhando no R": "Programação (R / Python)",
    "Primeiros com Python": "Programação (R / Python)",
    
    # Estatística & Probabilidade
    "Fundamentos da Estatística": "Estatística & Probabilidade",
    "Medidas de posição e dispersão": "Estatística & Probabilidade",
    "Distribuições e associações": "Estatística & Probabilidade",
    "Fundamentos de probabilidade": "Estatística & Probabilidade",
    "Probabilidade discreta e contínua": "Estatística & Probabilidade",
    
    # Visualização
    "Visualizações quantitativas: escolhendo o gráfico certo": "Visualização de Dados",
    "Visualização de dados Ggplot2": "Visualização de Dados",
    "Medidas estatísticas e visualização": "Visualização de Dados",
    
    # Fundamentos & ML
    "Big Data e Machine Learning": "ML & Fundamentos",
    "Pensamento crítico e analítico": "ML & Fundamentos",
    "Conceituação do curso e foguetes": "ML & Fundamentos",
    "Perguntas, contexto e organização da análise": "ML & Fundamentos"
}

def clean_str(s):
    if s is None:
        return ""
    return str(s).strip()

def time_to_seconds(t):
    if t is None:
        return 0
    if isinstance(t, datetime.time):
        return t.hour * 3600 + t.minute * 60 + t.second + (t.microsecond / 1_000_000)
    if isinstance(t, datetime.timedelta):
        return t.total_seconds()
    if isinstance(t, (int, float)):
        # If excel represented time as fraction of day
        return t * 86400
    try:
        parts = str(t).split(":")
        if len(parts) >= 2:
            h = float(parts[0])
            m = float(parts[1])
            s = float(parts[2]) if len(parts) > 2 else 0
            return h * 3600 + m * 60 + s
    except Exception:
        pass
    return 0

def get_periodo_dia(time_val):
    if not time_val:
        return "Indefinido"
    hour = 0
    if isinstance(time_val, datetime.time):
        hour = time_val.hour
    elif isinstance(time_val, str) and ":" in time_val:
        try:
            hour = int(time_val.split(":")[0])
        except:
            pass
    if 5 <= hour < 12:
        return "Manhã"
    elif 12 <= hour < 18:
        return "Tarde"
    else:
        return "Noite"

def get_dia_semana_pt(dt):
    dias = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]
    return dias[dt.weekday()]

records = []

for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        continue
        
    for r_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        disc = clean_str(row[0])
        dia_val = row[1]
        h_ini = row[2]
        h_fim = row[3]
        tempo_val = row[4]
        
        if not disc and not dia_val:
            continue
            
        dt = None
        if isinstance(dia_val, datetime.datetime):
            dt = dia_val.date()
        elif isinstance(dia_val, datetime.date):
            dt = dia_val
        elif isinstance(dia_val, str):
            try:
                dt = datetime.datetime.strptime(dia_val[:10], "%Y-%m-%d").date()
            except:
                pass
                
        segundos = time_to_seconds(tempo_val)
        
        # If duration was 0 or missing, try h_fim - h_ini
        if segundos <= 0 and h_ini and h_fim:
            s_ini = time_to_seconds(h_ini)
            s_fim = time_to_seconds(h_fim)
            if s_fim >= s_ini:
                segundos = s_fim - s_ini
                
        horas_dec = round(segundos / 3600.0, 2)
        minutos_totais = round(segundos / 60.0, 1)
        
        # Format string "Xh Ymin"
        h = int(segundos // 3600)
        m = int((segundos % 3600) // 60)
        duracao_str = f"{h}h {m:02d}m" if h > 0 else f"{m}m"
        
        hora_ini_str = h_ini.strftime("%H:%M:%S") if isinstance(h_ini, datetime.time) else (str(h_ini) if h_ini else "")
        hora_fim_str = h_fim.strftime("%H:%M:%S") if isinstance(h_fim, datetime.time) else (str(h_fim) if h_fim else "")
        
        periodo_dia = get_periodo_dia(h_ini)
        categoria = CATEGORIAS.get(disc, "Outros")
        
        records.append({
            "id": len(records) + 1,
            "fase": sheetname,
            "disciplina": disc,
            "categoria": categoria,
            "data": dt.strftime("%Y-%m-%d") if dt else "Indefinida",
            "data_formatada": dt.strftime("%d/%m/%Y") if dt else "Indefinida",
            "ano": dt.year if dt else 0,
            "mes": dt.month if dt else 0,
            "mes_ano": dt.strftime("%m/%Y") if dt else "",
            "dia_semana": get_dia_semana_pt(dt) if dt else "",
            "dia_semana_num": dt.weekday() if dt else -1,
            "hora_inicio": hora_ini_str,
            "hora_fim": hora_fim_str,
            "segundos": segundos,
            "minutos": minutos_totais,
            "horas": horas_dec,
            "duracao_formatada": duracao_str,
            "turno": periodo_dia
        })

# Sort records by date and hora_inicio
records.sort(key=lambda r: (r["data"], r["hora_inicio"]))

# Calculate cumulative hours
cum_h = 0.0
for r in records:
    cum_h += r["horas"]
    r["horas_acumuladas"] = round(cum_h, 2)

# Global metrics
total_segundos = sum(r["segundos"] for r in records)
total_horas = round(total_segundos / 3600.0, 1)
total_sessoes = len(records)
dias_unicos = len(set(r["data"] for r in records if r["data"] != "Indefinida"))
media_minutos_sessao = round((total_segundos / 60.0) / total_sessoes, 1) if total_sessoes > 0 else 0
media_horas_dia_estudado = round(total_horas / dias_unicos, 2) if dias_unicos > 0 else 0

# Timestamp e metadados de atualização com fuso horário de Brasília (UTC-3)
tz_brasilia = datetime.timezone(datetime.timedelta(hours=-3))
now = datetime.datetime.now(tz_brasilia)
data_atualizacao = now.strftime("%d/%m/%Y às %H:%M")
data_atualizacao_iso = now.isoformat()

try:
    mtime = os.path.getmtime(file_path)
    data_modificacao_planilha = datetime.datetime.fromtimestamp(mtime, tz=tz_brasilia).strftime("%d/%m/%Y às %H:%M")
except Exception:
    data_modificacao_planilha = data_atualizacao

ultima_sessao_formatada = records[-1]["data_formatada"] if records else ""

output_data = {
    "summary": {
        "total_horas": total_horas,
        "total_sessoes": total_sessoes,
        "dias_unicos_estudo": dias_unicos,
        "media_minutos_sessao": media_minutos_sessao,
        "media_horas_dia_estudado": media_horas_dia_estudado,
        "total_disciplinas": len(set(r["disciplina"] for r in records)),
        "data_inicio": records[0]["data_formatada"] if records else "",
        "data_fim": records[-1]["data_formatada"] if records else "",
        "data_atualizacao": data_atualizacao,
        "data_atualizacao_iso": data_atualizacao_iso,
        "data_modificacao_planilha": data_modificacao_planilha,
        "ultima_sessao": ultima_sessao_formatada
    },
    "records": records
}

# Save JSON
json_path = os.path.join(os.path.dirname(__file__), "study_data.json")
with open(json_path, "w", encoding="utf-8") as f:
    json.dump(output_data, f, ensure_ascii=False, indent=2)

# Save JS file for zero-CORS static standalone loading
js_path = os.path.join(os.path.dirname(__file__), "study_data.js")
with open(js_path, "w", encoding="utf-8") as f:
    f.write("window.STUDY_DATA = " + json.dumps(output_data, ensure_ascii=False, indent=2) + ";\n")

print(f"ETL finalizado com sucesso!")
print(f"Total Horas: {total_horas}h em {total_sessoes} sessões ({dias_unicos} dias únicos de estudo).")
print(f"Período: de {records[0]['data_formatada']} até {records[-1]['data_formatada']}")
print(f"Última atualização da base: {data_atualizacao}")

